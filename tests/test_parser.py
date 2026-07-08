"""Unit tests for config parser."""

import pytest
from src.parser import (
    parse_device_config, read_text_file, extract_text_from_rtf,
    extract_text_from_xlsx, extract_text_from_xls,
    extract_text_from_docx, extract_text_from_file,
)
from src.eol import lookup_eol


SAMPLE_CISCO_CONFIG = """
hostname CORE-SW-5510_M1
!
interface Vlan100
 ip address 10.100.1.1 255.255.255.0
!
interface GigabitEthernet1/0/1
 description Uplink
!
router ospf 1
 router-id 10.100.1.1
 network 10.100.0.0 0.0.255.255 area 0
!
ip route 0.0.0.0 0.0.0.0 10.100.1.254
!
snmp-server community public ro
!
access-list 100 permit ip any any
!
ntp server 10.100.0.1
!
"""

SAMPLE_HUAWEI_CONFIG = """
sysname hw-S6730-CORE
!
interface Vlanif100
 ip address 10.200.1.1 255.255.255.0
!
ospf 1 router-id 10.200.1.1
 area 0.0.0.0
  network 10.200.0.0 0.0.255.255
!
ip route-static 0.0.0.0 0.0.0.0 10.200.1.254
!
snmp-agent community read public
!
acl number 2000
 rule 5 permit source 10.0.0.0 0.255.255.255
!
ntp server 10.200.0.1
!
"""

SAMPLE_HP_CONFIG = """
hostname "HP-5412R-CORE"
!
interface 1
 name "Uplink"
 ip address 10.0.0.2 255.255.255.0
!
ip route 0.0.0.0 0.0.0.0 10.0.0.1
!
snmp-server community public
!
"""


class TestParser:
    def test_parse_cisco_config(self):
        device = parse_device_config(SAMPLE_CISCO_CONFIG, "cisco_config.txt")
        assert device.hostname == "CORE-SW-5510_M1"
        assert "5510" in device.model
        assert device.ip_mgmt == "10.100.1.1"
        assert len(device.vlans) > 0
        assert len(device.routes) > 0
        assert len(device.acl) > 0
        assert len(device.ntp_servers) > 0
        assert device.ospf.get("process_id") == "1"

    def test_parse_huawei_config(self):
        device = parse_device_config(SAMPLE_HUAWEI_CONFIG, "huawei_config.txt")
        assert device.hostname == "hw-S6730-CORE"
        assert "S6730" in device.model
        assert device.ip_mgmt == "10.200.1.1"
        assert device.ospf.get("router_id") == "10.200.1.1"

    def test_parse_hp_config(self):
        device = parse_device_config(SAMPLE_HP_CONFIG, "hp_config.txt")
        assert device.hostname == "HP-5412R-CORE"
        assert "5412" in device.model

    def test_empty_config(self):
        device = parse_device_config("", "empty.txt")
        assert device.hostname == ""
        assert device.model == ""

    def test_lookup_eol_found(self):
        result = lookup_eol(hostname="CORE-SW-5510_M1", model="HPE 5510 24G SFP 4SFP+ HI")
        assert result["status"] == "End-of-Sale"

    def test_lookup_eol_not_found(self):
        result = lookup_eol(hostname="unknown-device")
        assert result["status"] == "Требуется проверка"

    def test_read_text_file_utf8(self, tmp_path):
        f = tmp_path / "test.txt"
        f.write_text("hello world", encoding="utf-8")
        assert read_text_file(f) == "hello world"

    def test_extract_text_from_rtf(self, tmp_path):
        f = tmp_path / "test.rtf"
        f.write_text(r"{\rtf1\ansi\b Hello\b0 World}")
        result = extract_text_from_rtf(f)
        assert "Hello" in result
        assert "World" in result

    def test_parse_device_with_by_core_hostname(self):
        text = """
hostname BY_Core
!
interface Vlan1
 ip address 10.0.0.1 255.255.255.0
!
"""
        device = parse_device_config(text, "by_core.txt")
        assert device.hostname == "BY_Core"
        assert device.model == "Cisco ISR4331/K9" or device.model != ""

    def test_extract_text_from_xlsx(self, tmp_path):
        f = tmp_path / "test.xlsx"
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            ws.append(["Name", "Value"])
            ws.append(["Test", "123"])
            wb.save(str(f))
            result = extract_text_from_xlsx(f)
            assert "Name" in result
            assert "Test" in result
            assert "123" in result
        except ImportError:
            pytest.skip("openpyxl not available")

    def test_extract_text_from_xls(self, tmp_path):
        f = tmp_path / "test.xls"
        try:
            import xlwt
            wb = xlwt.Workbook()
            ws = wb.add_sheet("Sheet1")
            ws.write(0, 0, "Name")
            ws.write(0, 1, "Value")
            ws.write(1, 0, "Test")
            ws.write(1, 1, 456)
            wb.save(str(f))
            result = extract_text_from_xls(f)
            assert "Name" in result
            assert "Test" in result
        except ImportError:
            try:
                import pandas as pd
                df = pd.DataFrame({"Name": ["Test"], "Value": [456]})
                df.to_excel(str(f), index=False)
                result = extract_text_from_xls(f)
                assert "Name" in result or "Test" in result
            except ImportError:
                pytest.skip("neither xlwt nor pandas available")

    def test_extract_text_from_docx(self, tmp_path):
        f = tmp_path / "test.docx"
        try:
            from docx import Document
            doc = Document()
            doc.add_paragraph("Hello from docx")
            doc.save(str(f))
            result = extract_text_from_docx(f)
            assert "Hello from docx" in result
        except ImportError:
            pytest.skip("python-docx not available")

    def test_extract_text_from_file_dispatcher(self, tmp_path):
        f_txt = tmp_path / "test.txt"
        f_txt.write_text("plain text content", encoding="utf-8")
        result = extract_text_from_file(f_txt)
        assert result == "plain text content"

        f_csv = tmp_path / "test.csv"
        f_csv.write_text("a,b,c\n1,2,3", encoding="utf-8")
        result = extract_text_from_file(f_csv)
        assert "a,b,c" in result

    def test_extract_text_from_file_unknown_extension(self, tmp_path):
        f = tmp_path / "test.unknown"
        f.write_text("some content")
        result = extract_text_from_file(f)
        assert result == ""

    def test_template_section_extraction_numbered(self):
        from src.report import _extract_template_sections
        tmpl = """
1. ПЕРВЫЙ РАЗДЕЛ
Содержимое первого раздела

2. ВТОРОЙ РАЗДЕЛ
Содержимое второго раздела
"""
        sections = _extract_template_sections(tmpl)
        assert len(sections) >= 2
        assert sections[0]["title"] == "ПЕРВЫЙ РАЗДЕЛ"

    def test_template_section_extraction_uppercase(self):
        from src.report import _extract_template_sections
        tmpl = """
ВВЕДЕНИЕ
Текст введения.

ОСНОВНАЯ ЧАСТЬ
Текст основной части.
"""
        sections = _extract_template_sections(tmpl)
        assert len(sections) >= 2
        assert sections[0]["title"] == "ВВЕДЕНИЕ"

    def test_template_section_extraction(self):
        from src.report import _extract_template_sections
        tmpl = """
# Главный раздел

## Подраздел 1

Содержимое 1

## Подраздел 2

Содержимое 2
"""
        sections = _extract_template_sections(tmpl)
        assert len(sections) >= 1
        assert sections[0]["title"] == "Главный раздел"
        has_sub = any(s["title"] == "Подраздел 1" for s in sections)
        assert has_sub


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
